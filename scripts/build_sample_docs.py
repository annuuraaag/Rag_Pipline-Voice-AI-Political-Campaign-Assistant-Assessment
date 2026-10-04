"""Render the fictional sample corpus from Markdown sources into PDF / DOCX / MD / TXT.

Sources live in sample_data/_source/**.md with a front-matter `output:` path whose
extension decides the format. PDFs get running footers and page breaks (to exercise
header/footer removal and page-level citations); DOCX files use real heading styles.

    python scripts/build_sample_docs.py
"""
from __future__ import annotations

import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "sample_data" / "_source"
OUT = ROOT / "sample_data"

# Corpus manifest: the metadata an operator would attach at upload time.
METADATA = {
    "manifesto/sca_manifesto_2026.pdf": dict(district="statewide", category="manifesto", topic="general", source="SCA People's Manifesto 2026"),
    "districts/vijayawada_district_plan.pdf": dict(district="vijayawada", category="district_profile", topic="general", source="SCA District Plans 2026"),
    "districts/guntur_district_plan.md": dict(district="guntur", category="district_profile", topic="general", source="SCA District Plans 2026"),
    "districts/visakhapatnam_district_plan.docx": dict(district="visakhapatnam", category="district_profile", topic="general", source="SCA District Plans 2026"),
    "candidate/candidate_profile.docx": dict(district="vijayawada", category="candidate_profile", topic="general", source="SCA Candidate Profiles", candidate="Dr. Anitha Rao Kotagiri"),
    "schemes/healthcare_schemes.md": dict(district="statewide", category="scheme", topic="healthcare", source="SCA Scheme Guides 2026"),
    "schemes/education_schemes.txt": dict(district="statewide", category="scheme", topic="education", source="SCA Scheme Guides 2026"),
    "employment/employment_plan.md": dict(district="statewide", category="policy", topic="employment", source="SCA Policy Papers 2026"),
    "faq/campaign_faq.md": dict(district="statewide", category="faq", topic="campaign_info", source="SCA Campaign FAQ"),
    "agriculture/farmers_charter.txt": dict(district="statewide", category="policy", topic="agriculture", source="SCA Policy Papers 2026"),
}


@dataclass
class Node:
    kind: str            # heading | para | bullet | table | pagebreak
    text: str = ""
    level: int = 0
    rows: list[list[str]] = field(default_factory=list)


@dataclass
class SourceDoc:
    front: dict[str, str]
    title: str
    nodes: list[Node]
    raw_body: str


def parse_source(path: Path) -> SourceDoc:
    text = path.read_text(encoding="utf-8")
    front: dict[str, str] = {}
    m = re.match(r"\A---\n(.*?)\n---\n", text, re.DOTALL)
    if m:
        for line in m.group(1).splitlines():
            k, _, v = line.partition(":")
            front[k.strip()] = v.strip()
        text = text[m.end():]
    title, nodes, para, table = "", [], [], []

    def flush():
        nonlocal para, table
        if para:
            nodes.append(Node("para", " ".join(para)))
        if table:
            nodes.append(Node("table", rows=table))
        para, table = [], []

    for line in text.splitlines():
        s = line.strip()
        if s == "<!-- pagebreak -->":
            flush(); nodes.append(Node("pagebreak")); continue
        if not s:
            flush(); continue
        h = re.match(r"^(#{1,6})\s+(.*)$", s)
        if h:
            flush()
            if len(h.group(1)) == 1:
                title = h.group(2)
            else:
                nodes.append(Node("heading", h.group(2), level=len(h.group(1)) - 1))
            continue
        if s.startswith("|"):
            if para:
                nodes.append(Node("para", " ".join(para))); para = []
            cells = [c.strip() for c in s.strip("|").split("|")]
            if not all(re.fullmatch(r"-{3,}", c) for c in cells):
                table.append(cells)
            continue
        if s.startswith("- "):
            flush(); nodes.append(Node("bullet", s[2:])); continue
        para.append(s)
    flush()
    return SourceDoc(front, title, nodes, text)


def _plain(t: str) -> str:
    return re.sub(r"\*\*(.*?)\*\*", r"\1", t)


# ── renderers ────────────────────────────────────────────────────────────
def render_md(doc: SourceDoc, out: Path) -> None:
    meta = {k: v for k, v in doc.front.items() if k not in ("output", "footer")}
    fm = ("---\n" + "\n".join(f"{k}: {v}" for k, v in meta.items()) + "\n---\n") if meta else ""
    body = doc.raw_body.replace("<!-- pagebreak -->\n\n", "")
    out.write_text(fm + body.lstrip(), encoding="utf-8")


def render_txt(doc: SourceDoc, out: Path) -> None:
    lines = [doc.title, ""]
    for n in doc.nodes:
        if n.kind == "heading":
            lines += [n.text, ""]
        elif n.kind == "para":
            lines += [_plain(n.text), ""]
        elif n.kind == "bullet":
            lines.append(f"- {_plain(n.text)}")
        elif n.kind == "table":
            lines += [" | ".join(r) for r in n.rows] + [""]
    out.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def render_docx(doc: SourceDoc, out: Path) -> None:
    import docx

    d = docx.Document()
    d.core_properties.title = doc.title
    d.add_heading(doc.title, 0)
    for n in doc.nodes:
        if n.kind == "heading":
            d.add_heading(n.text, n.level)
        elif n.kind == "para":
            d.add_paragraph(_plain(n.text))
        elif n.kind == "bullet":
            d.add_paragraph(_plain(n.text), style="List Bullet")
        elif n.kind == "table":
            t = d.add_table(rows=len(n.rows), cols=len(n.rows[0]))
            t.style = "Table Grid"
            for i, row in enumerate(n.rows):
                for j, cell in enumerate(row):
                    t.cell(i, j).text = cell
    d.save(out)


def render_pdf(doc: SourceDoc, out: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table

    styles = getSampleStyleSheet()
    footer = doc.front.get("footer", doc.title)

    def on_page(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(2 * cm, 1.2 * cm, f"{footer} · Page {d.page}")
        canvas.restoreState()

    story = [Paragraph(doc.title, styles["Title"]), Spacer(1, 6)]
    for n in doc.nodes:
        if n.kind == "heading":
            story.append(Paragraph(n.text, styles["Heading2" if n.level == 1 else "Heading3"]))
        elif n.kind == "para":
            story.append(Paragraph(re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", n.text), styles["BodyText"]))
        elif n.kind == "bullet":
            story.append(Paragraph(n.text, styles["BodyText"], bulletText="•"))
        elif n.kind == "table":
            story.append(Table(n.rows))
        elif n.kind == "pagebreak":
            story.append(PageBreak())
    SimpleDocTemplate(str(out), pagesize=A4, title=doc.title, author="Sunrise Coast Alliance (fictional)",
                      leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm
                      ).build(story, onFirstPage=on_page, onLaterPages=on_page)


def build_scanned_samples(out_dir: Path) -> None:
    """Image-only documents (no text layer) to demonstrate the OCR fallback. Not in the seeded corpus."""
    from PIL import Image, ImageDraw, ImageFont

    fonts = Path("/usr/share/fonts/truetype/dejavu")
    body, bold = ImageFont.truetype(str(fonts / "DejaVuSans.ttf"), 30), ImageFont.truetype(str(fonts / "DejaVuSans-Bold.ttf"), 40)
    head = ImageFont.truetype(str(fonts / "DejaVuSans-Bold.ttf"), 32)
    lines = [
        ("Kurnool Water Security Plan", bold), ("", body),
        ("FICTIONAL SAMPLE DOCUMENT. Scanned copy for an OCR demo; all figures are invented.", body), ("", body),
        ("1. Drinking Water", head),
        ("The Sunrise Coast Alliance will lay a 140 kilometre pipeline from the", body),
        ("Tungabhadra river to supply drinking water to 220 villages in Kurnool", body),
        ("district. Every village will get a water testing kiosk by 2028.", body), ("", body),
        ("2. Irrigation", head),
        ("Farmers in Kurnool will receive drip irrigation kits at a 90 percent", body),
        ("subsidy, and 35 check dams will be built on seasonal streams.", body),
    ]
    img = Image.new("RGB", (1654, 2339), "white")   # A4 at 200 dpi
    draw = ImageDraw.Draw(img)
    y = 180
    for text, font in lines:
        draw.text((150, y), text, fill="black", font=font)
        y += 64 if text else 40
    draw.text((150, 2240), "SCA District Plans - Kurnool - Page 1", fill="black", font=body)
    out_dir.mkdir(parents=True, exist_ok=True)
    img.save(out_dir / "kurnool_water_plan_scanned.pdf", resolution=200)
    img.save(out_dir / "kurnool_water_plan_photo.png")
    print("built extra/kurnool_water_plan_scanned.pdf and extra/kurnool_water_plan_photo.png (no text layer)")


RENDERERS = {".md": render_md, ".txt": render_txt, ".docx": render_docx, ".pdf": render_pdf}


def main() -> int:
    manifest = {"description": "Fictional campaign corpus for the Campaign Voice RAG assessment. "
                               "All parties, people, schemes and figures are invented.", "documents": []}
    for src in sorted(SRC.rglob("*.md")):
        doc = parse_source(src)
        rel = doc.front["output"]
        out = OUT / rel
        out.parent.mkdir(parents=True, exist_ok=True)
        RENDERERS[out.suffix](doc, out)
        manifest["documents"].append({"path": rel, **METADATA[rel]})
        print(f"built {rel} ({out.stat().st_size:,} bytes)")
    missing = set(METADATA) - {d["path"] for d in manifest["documents"]}
    if missing:
        print(f"ERROR: manifest entries without sources: {missing}", file=sys.stderr)
        return 1
    manifest["documents"].sort(key=lambda d: d["path"])
    build_scanned_samples(OUT / "extra")
    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(f"wrote manifest with {len(manifest['documents'])} documents")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
